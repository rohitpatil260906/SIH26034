import os
import io
from typing import Dict, Any, List, Optional
from datetime import datetime

from reportlab.lib.pagesizes import A4
from reportlab.lib import colors
from reportlab.lib.styles import getSampleStyleSheet, ParagraphStyle
from reportlab.platypus import (
    SimpleDocTemplate, Paragraph, Spacer, Table, TableStyle, PageBreak, KeepTogether, HRFlowable
)
from reportlab.lib.units import inch

from docx import Document
from docx.shared import Inches, Pt, RGBColor
from docx.enum.text import WD_ALIGN_PARAGRAPH

from ..models import ScanProcessResponse, ComplianceCheckItem, CanonicalField

REPORTS_DIR = os.path.join(os.path.dirname(os.path.dirname(__file__)), "data", "reports")
os.makedirs(REPORTS_DIR, exist_ok=True)

# -------------------------------------------------------------------
# STAGE 4: REPORTLAB PDF GENERATION
# -------------------------------------------------------------------

def generate_statutory_pdf_report(scan: ScanProcessResponse) -> str:
    """Generates the official Department of Consumer Affairs Legal Metrology PDF Report."""
    report_filename = f"report_{scan.scan_id}.pdf"
    report_path = os.path.join(REPORTS_DIR, report_filename)
    
    doc = SimpleDocTemplate(
        report_path,
        pagesize=A4,
        rightMargin=36,
        leftMargin=36,
        topMargin=36,
        bottomMargin=36
    )
    
    styles = getSampleStyleSheet()
    
    # Custom styles
    title_style = ParagraphStyle(
        'GovtTitle',
        parent=styles['Normal'],
        fontName='Helvetica-Bold',
        fontSize=13,
        leading=16,
        alignment=1, # Center
        textColor=colors.HexColor('#0f2942')
    )
    
    subtitle_style = ParagraphStyle(
        'GovtSubTitle',
        parent=styles['Normal'],
        fontName='Helvetica-Bold',
        fontSize=10,
        leading=13,
        alignment=1,
        textColor=colors.HexColor('#334155')
    )

    report_name_style = ParagraphStyle(
        'ReportName',
        parent=styles['Normal'],
        fontName='Helvetica-Bold',
        fontSize=12,
        leading=15,
        alignment=1,
        textColor=colors.HexColor('#991b1b')
    )
    
    heading2_style = ParagraphStyle(
        'SectionH2',
        parent=styles['Heading2'],
        fontName='Helvetica-Bold',
        fontSize=10,
        leading=13,
        textColor=colors.HexColor('#0f2942'),
        spaceBefore=10,
        spaceAfter=4
    )
    
    cell_style = ParagraphStyle(
        'TableCell',
        parent=styles['Normal'],
        fontName='Helvetica',
        fontSize=8,
        leading=10,
        textColor=colors.HexColor('#1e293b')
    )

    cell_bold_style = ParagraphStyle(
        'TableCellBold',
        parent=styles['Normal'],
        fontName='Helvetica-Bold',
        fontSize=8,
        leading=10,
        textColor=colors.HexColor('#0f2942')
    )

    story = []
    
    # Header
    story.append(Paragraph("GOVERNMENT OF INDIA", title_style))
    story.append(Paragraph("MINISTRY OF CONSUMER AFFAIRS, FOOD & PUBLIC DISTRIBUTION", subtitle_style))
    story.append(Paragraph("DEPARTMENT OF CONSUMER AFFAIRS • LEGAL METROLOGY DIVISION", subtitle_style))
    story.append(Spacer(1, 4))
    story.append(Paragraph("PACKAGED COMMODITY COMPLIANCE INSPECTION REPORT", report_name_style))
    story.append(Paragraph("Statutory Audit under The Legal Metrology Act, 2009 & Packaged Commodities Rules, 2011", subtitle_style))
    story.append(Spacer(1, 8))
    story.append(HRFlowable(width="100%", thickness=1.5, color=colors.HexColor('#0f2942'), spaceAfter=8))
    
    # Particulars Table
    infractions = [c for c in scan.compliance_checks if c.status == "FAIL"]
    reviews = [c for c in scan.compliance_checks if c.status == "NEEDS REVIEW"]
    
    pdata = [
        [
            Paragraph("<b>Case / Docket ID:</b>", cell_style), Paragraph(scan.scan_id, cell_bold_style),
            Paragraph("<b>Inspection Date:</b>", cell_style), Paragraph(scan.timestamp[:19].replace("T", " "), cell_style)
        ],
        [
            Paragraph("<b>Inspecting Officer:</b>", cell_style), Paragraph("Inspector LM-0842", cell_style),
            Paragraph("<b>Compliance Score:</b>", cell_style), Paragraph(f"<b>{scan.compliance_score} / 100</b> ({scan.overall_status})", cell_bold_style)
        ],
        [
            Paragraph("<b>Product Name:</b>", cell_style), Paragraph(scan.product_info.product_name, cell_bold_style),
            Paragraph("<b>Commodity Type:</b>", cell_style), Paragraph(scan.product_info.commodity_name, cell_style)
        ],
        [
            Paragraph("<b>Manufacturer:</b>", cell_style), Paragraph(scan.product_info.manufacturer.full_address or "Not declared", cell_style),
            Paragraph("<b>Statutory Violations:</b>", cell_style), Paragraph(f"<font color='{'#b91c1c' if infractions else '#15803d'}'><b>{len(infractions)} active infractions</b></font>", cell_bold_style)
        ],
        [
            Paragraph("<b>Surfaces Audited:</b>", cell_style), Paragraph(", ".join(scan.surfaces_processed), cell_style),
            Paragraph("<b>External Verification:</b>", cell_style), Paragraph(f"<i>{scan.external_verification}</i>", cell_style)
        ]
    ]
    
    ptable = Table(pdata, colWidths=[1.3*inch, 2.3*inch, 1.4*inch, 2.2*inch])
    ptable.setStyle(TableStyle([
        ('BACKGROUND', (0, 0), (-1, -1), colors.HexColor('#f8fafc')),
        ('BOX', (0, 0), (-1, -1), 1, colors.HexColor('#cbd5e1')),
        ('INNERGRID', (0, 0), (-1, -1), 0.5, colors.HexColor('#e2e8f0')),
        ('TOPPADDING', (0, 0), (-1, -1), 3),
        ('BOTTOMPADDING', (0, 0), (-1, -1), 3),
    ]))
    story.append(ptable)
    story.append(Spacer(1, 10))

    # SECTION 1: EXTRACTED INFORMATION
    story.append(Paragraph("1. EXTRACTED STATUTORY DECLARATIONS & CONFIDENCE MATRIX", heading2_style))
    
    ext_data = [
        [
            Paragraph("<b>Statutory Declaration Field</b>", cell_bold_style),
            Paragraph("<b>Detected Value</b>", cell_bold_style),
            Paragraph("<b>Confidence</b>", cell_bold_style),
            Paragraph("<b>Rule Reference</b>", cell_bold_style),
            Paragraph("<b>Status</b>", cell_bold_style)
        ]
    ]
    
    for cf in scan.canonical_fields:
        status_color = "#15803d" if cf.status == "Found" else ("#b91c1c" if cf.status == "Defective" else "#d97706")
        ext_data.append([
            Paragraph(cf.statutory_name, cell_style),
            Paragraph(cf.extracted_value, cell_style),
            Paragraph(f"{int(cf.confidence * 100)}%", cell_style),
            Paragraph(cf.rule_reference, cell_style),
            Paragraph(f"<font color='{status_color}'><b>{cf.status}</b></font>", cell_style)
        ])
        
    ext_table = Table(ext_data, colWidths=[2.2*inch, 2.4*inch, 0.7*inch, 1.1*inch, 0.8*inch])
    ext_table.setStyle(TableStyle([
        ('BACKGROUND', (0, 0), (-1, 0), colors.HexColor('#f1f5f9')),
        ('BOX', (0, 0), (-1, -1), 1, colors.HexColor('#cbd5e1')),
        ('INNERGRID', (0, 0), (-1, -1), 0.5, colors.HexColor('#e2e8f0')),
        ('TOPPADDING', (0, 0), (-1, -1), 3),
        ('BOTTOMPADDING', (0, 0), (-1, -1), 3),
    ]))
    story.append(ext_table)
    story.append(Spacer(1, 10))

    # SECTION 2: COMPLIANCE CHECK
    story.append(Paragraph("2. STATUTORY LEGAL METROLOGY RULES COMPLIANCE AUDIT", heading2_style))
    
    chk_data = [
        [
            Paragraph("<b>Rule</b>", cell_bold_style),
            Paragraph("<b>Requirement & Sub-Rule</b>", cell_bold_style),
            Paragraph("<b>Audited Finding / Declaration</b>", cell_bold_style),
            Paragraph("<b>Verdict</b>", cell_bold_style)
        ]
    ]
    
    for chk in scan.compliance_checks[:14]:  # Primary critical rules on page
        v_color = "#15803d" if chk.status == "PASS" else ("#b91c1c" if chk.status == "FAIL" else "#d97706")
        chk_data.append([
            Paragraph(chk.rule_no, cell_bold_style),
            Paragraph(f"<b>{chk.rule_title}</b><br/>{chk.sub_rule}", cell_style),
            Paragraph(chk.detected_declaration, cell_style),
            Paragraph(f"<font color='{v_color}'><b>{chk.status}</b></font>", cell_style)
        ])
        
    chk_table = Table(chk_data, colWidths=[1.1*inch, 2.5*inch, 2.7*inch, 0.9*inch])
    chk_table.setStyle(TableStyle([
        ('BACKGROUND', (0, 0), (-1, 0), colors.HexColor('#f1f5f9')),
        ('BOX', (0, 0), (-1, -1), 1, colors.HexColor('#cbd5e1')),
        ('INNERGRID', (0, 0), (-1, -1), 0.5, colors.HexColor('#e2e8f0')),
        ('TOPPADDING', (0, 0), (-1, -1), 3),
        ('BOTTOMPADDING', (0, 0), (-1, -1), 3),
    ]))
    story.append(chk_table)
    story.append(Spacer(1, 10))

    # SECTION 3: VIOLATIONS & INFRACTIONS
    if infractions:
        story.append(Paragraph("3. STATUTORY VIOLATIONS & PENAL ACTION RECOMMENDATIONS", heading2_style))
        vio_data = [
            [
                Paragraph("<b>Infraction Clause</b>", cell_bold_style),
                Paragraph("<b>Detected Packaging Finding</b>", cell_bold_style),
                Paragraph("<b>Statutory Penalty / Compounding</b>", cell_bold_style)
            ]
        ]
        for inf in infractions:
            vio_data.append([
                Paragraph(f"<b>{inf.rule_no}</b> ({inf.rule_title})", cell_bold_style),
                Paragraph(inf.detected_declaration, cell_style),
                Paragraph(inf.section_penalty or "Section 36(1) Compounding fee: ₹25,000", cell_style)
            ])
        vio_table = Table(vio_data, colWidths=[2.2*inch, 2.8*inch, 2.2*inch])
        vio_table.setStyle(TableStyle([
            ('BACKGROUND', (0, 0), (-1, 0), colors.HexColor('#fee2e2')),
            ('BOX', (0, 0), (-1, -1), 1, colors.HexColor('#f87171')),
            ('INNERGRID', (0, 0), (-1, -1), 0.5, colors.HexColor('#fecaca')),
            ('TOPPADDING', (0, 0), (-1, -1), 3),
            ('BOTTOMPADDING', (0, 0), (-1, -1), 3),
        ]))
        story.append(vio_table)
        story.append(Spacer(1, 10))

    # SECTION 4: NEEDS REVIEW
    if reviews:
        story.append(Paragraph("4. INSPECTION ADVISORY — AMBIGUOUS FIELDS FLAGGED FOR HUMAN REVIEW", heading2_style))
        rev_data = [
            [
                Paragraph("<b>Flagged Field</b>", cell_bold_style),
                Paragraph("<b>Observation & Review Rationale</b>", cell_bold_style),
                Paragraph("<b>Statutory Rule</b>", cell_bold_style)
            ]
        ]
        for r in reviews:
            rev_data.append([
                Paragraph(r.rule_no, cell_bold_style),
                Paragraph(r.detected_declaration, cell_style),
                Paragraph(r.sub_rule, cell_style)
            ])
        rev_table = Table(rev_data, colWidths=[1.8*inch, 3.8*inch, 1.6*inch])
        rev_table.setStyle(TableStyle([
            ('BACKGROUND', (0, 0), (-1, 0), colors.HexColor('#fef3c7')),
            ('BOX', (0, 0), (-1, -1), 1, colors.HexColor('#fde68a')),
            ('INNERGRID', (0, 0), (-1, -1), 0.5, colors.HexColor('#fef3c7')),
            ('TOPPADDING', (0, 0), (-1, -1), 3),
            ('BOTTOMPADDING', (0, 0), (-1, -1), 3),
        ]))
        story.append(rev_table)
        story.append(Spacer(1, 10))

    # Sign-off footer
    story.append(Spacer(1, 12))
    story.append(HRFlowable(width="100%", thickness=0.8, color=colors.HexColor('#94a3b8'), spaceAfter=8))
    
    footer_data = [
        [
            Paragraph("<b>Digital Verification QR Seal</b><br/>Hash: " + scan.scan_id, cell_style),
            Paragraph("<b>Legal Metrology Officer Attestation</b><br/>Authorized Enforcement Officer Signature", cell_style)
        ]
    ]
    ftable = Table(footer_data, colWidths=[4.0*inch, 3.2*inch])
    ftable.setStyle(TableStyle([
        ('ALIGN', (1, 0), (1, 0), 'RIGHT'),
    ]))
    story.append(ftable)

    doc.build(story)
    return report_path

# -------------------------------------------------------------------
# STAGE 4: DOCX REPORT GENERATION
# -------------------------------------------------------------------

def generate_statutory_docx_report(scan: ScanProcessResponse) -> str:
    """Generates the official Department of Consumer Affairs Legal Metrology Word Report."""
    report_filename = f"report_{scan.scan_id}.docx"
    report_path = os.path.join(REPORTS_DIR, report_filename)
    
    doc = Document()
    
    # Title
    p_title = doc.add_paragraph()
    p_title.alignment = WD_ALIGN_PARAGRAPH.CENTER
    run_t = p_title.add_run("DEPARTMENT OF CONSUMER AFFAIRS\nLEGAL METROLOGY DIVISION")
    run_t.bold = True
    run_t.font.size = Pt(14)
    run_t.font.color.rgb = RGBColor(15, 41, 66)

    p_sub = doc.add_paragraph()
    p_sub.alignment = WD_ALIGN_PARAGRAPH.CENTER
    run_s = p_sub.add_run("PACKAGED COMMODITY COMPLIANCE INSPECTION REPORT\nUnder Legal Metrology (Packaged Commodities) Rules, 2011\n")
    run_s.bold = True
    run_s.font.size = Pt(11)
    run_s.font.color.rgb = RGBColor(153, 27, 27)

    # Particulars
    doc.add_heading("1. Inspection Particulars", level=2)
    p_details = doc.add_paragraph()
    p_details.add_run(f"Case ID: {scan.scan_id}\n")
    p_details.add_run(f"Inspection Timestamp: {scan.timestamp}\n")
    p_details.add_run(f"Product Name: {scan.product_info.product_name}\n")
    p_details.add_run(f"Manufacturer: {scan.product_info.manufacturer.full_address}\n")
    p_details.add_run(f"Compliance Score: {scan.compliance_score} / 100 ({scan.overall_status})\n")
    p_details.add_run(f"External Verification: {scan.external_verification}\n")

    # Extracted Info Table
    doc.add_heading("2. Extracted Statutory Information", level=2)
    table = doc.add_table(rows=1, cols=4)
    hdr_cells = table.rows[0].cells
    hdr_cells[0].text = 'Field'
    hdr_cells[1].text = 'Detected Value'
    hdr_cells[2].text = 'Confidence'
    hdr_cells[3].text = 'Status'
    
    for cf in scan.canonical_fields:
        row_cells = table.add_row().cells
        row_cells[0].text = cf.statutory_name
        row_cells[1].text = cf.extracted_value
        row_cells[2].text = f"{int(cf.confidence * 100)}%"
        row_cells[3].text = cf.status

    # Violations
    infractions = [c for c in scan.compliance_checks if c.status == "FAIL"]
    if infractions:
        doc.add_heading("3. Statutory Violations & Compounding Directives", level=2)
        v_table = doc.add_table(rows=1, cols=3)
        v_cells = v_table.rows[0].cells
        v_cells[0].text = 'Rule'
        v_cells[1].text = 'Detected Infraction'
        v_cells[2].text = 'Penal Provision'
        
        for inf in infractions:
            v_row = v_table.add_row().cells
            v_row[0].text = f"{inf.rule_no} ({inf.rule_title})"
            v_row[1].text = inf.detected_declaration
            v_row[2].text = inf.section_penalty or "Section 36(1) Compounding fee: ₹25,000"

    doc.save(report_path)
    return report_path
